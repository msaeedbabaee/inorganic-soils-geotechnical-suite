import io
import numpy as np
import openpyxl
import pandas as pd
import plotly.graph_objects as go
from scipy.interpolate import interp1d
import streamlit as st

# --- PAGE CONFIGURATION ---
st.set_page_config(
    page_title="Inorganic Soils Suite", page_icon="🧪", layout="wide"
)


# --- OOP COMPUTATIONAL ENGINES ---
class GrainSizeAnalyzer:

    def __init__(self, df_sieve: pd.DataFrame):
        self.df = df_sieve.sort_values(
            by="grain_size_mm", ascending=False
        ).reset_index(drop=True)

    def calculate_coefficients(self):
        sizes = self.df["grain_size_mm"].values
        passing = self.df["percent_passing"].values
        valid = (sizes > 0) & (passing > 0) & (passing < 100)
        if sum(valid) < 2:
            return None, None, None, 0.0, 0.0, "Insufficient data points."
        try:
            log_sizes = np.log10(sizes[valid])
            pass_vals = passing[valid]
            f_interp = interp1d(
                pass_vals, log_sizes, kind="linear", fill_value="extrapolate"
            )
            d10 = 10 ** float(f_interp(10.0))
            d30 = 10 ** float(f_interp(30.0))
            d60 = 10 ** float(f_interp(60.0))
            cu = d60 / d10 if d10 > 0 else 0.0
            cc = (d30**2) / (d10 * d60) if (d10 * d60) > 0 else 0.0
            return d10, d30, d60, cu, cc, "Success"
        except Exception as e:
            return None, None, None, 0.0, 0.0, str(e)


class VisualSoilDescriptionEngine:

    def __init__(
        self,
        soil_name,
        color,
        moisture,
        plasticity,
        dry_strength,
        dilatancy,
        toughness,
        angularity,
        particle_shape,
        hcl_reaction,
    ):
        self.soil_name = soil_name
        self.color = color
        self.moisture = moisture
        self.plasticity = plasticity
        self.dry_strength = dry_strength
        self.dilatancy = dilatancy
        self.toughness = toughness
        self.angularity = angularity
        self.particle_shape = particle_shape
        self.hcl_reaction = hcl_reaction

    def generate_description_string(self) -> str:
        desc = f"{self.color}, {self.moisture} {self.soil_name}; "
        if self.soil_name in ["Lean Clay", "Fat Clay", "Elastic Silt", "Silt"]:
            desc += f"Plasticity: {self.plasticity}, Dry Strength: {self.dry_strength}, "
            desc += f"Dilatancy: {self.dilatancy}, Toughness: {self.toughness}."
        else:
            desc += f"Particle Angularity: {self.angularity}, Shape: {self.particle_shape}."
        if self.hcl_reaction != "None":
            desc += f" Reaction with HCl: {self.hcl_reaction}."
        return desc


class RelativeDensityCalculator:

    def __init__(
        self,
        method="Void Ratio",
        e=0.65,
        e_min=0.45,
        e_max=0.85,
        gamma_d=16.5,
        gamma_d_min=14.0,
        gamma_d_max=18.5,
    ):
        self.method = method
        self.e = e
        self.e_min = e_min
        self.e_max = e_max
        self.gamma_d = gamma_d
        self.gamma_d_min = gamma_d_min
        self.gamma_d_max = gamma_d_max

    def compute(self):
        if self.method.startswith("Void"):
            if self.e_max == self.e_min:
                return 0.0, "Invalid Input"
            dr = ((self.e_max - self.e) / (self.e_max - self.e_min)) * 100.0
        else:
            if self.gamma_d_max == self.gamma_d_min:
                return 0.0, "Invalid Input"
            dr = (
                (self.gamma_d_max * (self.gamma_d - self.gamma_d_min))
                / (self.gamma_d * (self.gamma_d_max - self.gamma_d_min))
            ) * 100.0

        dr = float(np.clip(dr, 0.0, 100.0))
        if dr < 15:
            compactness = "Very Loose"
        elif dr < 35:
            compactness = "Loose"
        elif dr < 65:
            compactness = "Medium Dense"
        elif dr < 85:
            compactness = "Dense"
        else:
            compactness = "Very Dense"
        return dr, compactness


# --- UI ---
st.title("🧪 Inorganic Soils Geotechnical Suite (CFEM Chapter 4)")

tab1, tab2, tab3 = st.tabs(
    [
        "📈 Grain Size & Gradation",
        "📝 Visual-Manual Description",
        "⚖️ Relative Density & Compactness",
    ]
)

with tab1:
    st.header("Grain Size Distribution & Gradation Curve Analyzer")
    col1, col2 = st.columns([1, 2])
    with col1:
        default_data = pd.DataFrame(
            {
                "grain_size_mm": [
                    50.0,
                    25.0,
                    9.5,
                    4.75,
                    2.0,
                    0.85,
                    0.425,
                    0.150,
                    0.075,
                ],
                "percent_passing": [
                    100.0,
                    90.0,
                    75.0,
                    60.0,
                    45.0,
                    32.0,
                    22.0,
                    10.0,
                    5.0,
                ],
            }
        )
        edited_sieve_df = st.data_editor(
            default_data, num_rows="dynamic", use_container_width=True
        )
    with col2:
        if not edited_sieve_df.empty:
            analyzer = GrainSizeAnalyzer(edited_sieve_df)
            d10, d30, d60, cu, cc, status = analyzer.calculate_coefficients()
            if status == "Success":
                m1, m2, m3, m4, m5 = st.columns(5)
                m1.metric("D10 (mm)", f"{d10:.3f}")
                m2.metric("D30 (mm)", f"{d30:.3f}")
                m3.metric("D60 (mm)", f"{d60:.3f}")
                m4.metric("Cu", f"{cu:.2f}")
                m5.metric("Cc", f"{cc:.2f}")

                fig = go.Figure()
                fig.add_trace(
                    go.Scatter(
                        x=edited_sieve_df["grain_size_mm"],
                        y=edited_sieve_df["percent_passing"],
                        mode="lines+markers",
                        name="Gradation Curve",
                    )
                )
                fig.update_xaxes(
                    type="log",
                    autorange="reversed",
                    title="Grain Size (mm) [Log Scale]",
                )
                fig.update_yaxes(title="Percent Passing (%)", range=[0, 100])
                fig.update_layout(
                    title="Particle Size Distribution Curve", height=350
                )
                st.plotly_chart(fig, use_container_width=True)

with tab2:
    st.header("Visual-Manual Soil Description Generator (ASTM D2488)")
    c1, c2, c3 = st.columns(3)
    with c1:
        soil_name = st.selectbox(
            "Primary Soil Type",
            [
                "Lean Clay",
                "Fat Clay",
                "Silt",
                "Elastic Silt",
                "Sand",
                "Gravel",
                "Sandy Lean Clay",
            ],
        )
        color = st.selectbox(
            "Color", ["Brown", "Dark Brown", "Gray", "Dark Gray"]
        )
        moisture = st.selectbox("Moisture Condition", ["Dry", "Moist", "Wet"])
    with c2:
        plasticity = st.selectbox(
            "Plasticity", ["Non-plastic", "Low", "Medium", "High"]
        )
        dry_strength = st.selectbox(
            "Dry Strength", ["None", "Low", "Medium", "High"]
        )
        dilatancy = st.selectbox("Dilatancy", ["None", "Slow", "Rapid"])
    with c3:
        toughness = st.selectbox("Toughness", ["Low", "Medium", "High"])
        angularity = st.selectbox(
            "Angularity", ["Angular", "Sub-angular", "Sub-rounded", "Rounded"]
        )
        particle_shape = st.selectbox("Particle Shape", ["Flat", "None"])
        hcl_reaction = st.selectbox(
            "Reaction with HCl", ["None", "Weak", "Strong"]
        )

    desc_engine = VisualSoilDescriptionEngine(
        soil_name,
        color,
        moisture,
        plasticity,
        dry_strength,
        dilatancy,
        toughness,
        angularity,
        particle_shape,
        hcl_reaction,
    )
    final_desc = desc_engine.generate_description_string()
    st.success(f"**Standard Descriptive Name:** \n\n `{final_desc}`")

with tab3:
    st.header("Relative Density & Compactness Calculator")
    method = st.radio(
        "Calculation Method:",
        ("Void Ratio (\(e, e_{max}, e_{min}\))", "Dry Unit Weight"),
    )

    if method.startswith("Void"):
        col_a, col_b, col_c = st.columns(3)
        with col_a:
            e = st.number_input("In-situ Void Ratio (\(e\))", value=0.65)
        with col_b:
            e_min = st.number_input("Min Void Ratio (\(e_{min}\))", value=0.45)
        with col_c:
            e_max = st.number_input("Max Void Ratio (\(e_{max}\))", value=0.85)
        rd_calc = RelativeDensityCalculator(
            method="Void Ratio", e=e, e_min=e_min, e_max=e_max
        )
    else:
        col_a, col_b, col_c = st.columns(3)
        with col_a:
            gamma_d = st.number_input(
                "Dry Unit Weight \(\\gamma_d\) (\(kN/m^3\))", value=16.5
            )
        with col_b:
            gamma_d_min = st.number_input(
                "Min \(\\gamma_d\) (\(kN/m^3\))", value=14.0
            )
        with col_c:
            gamma_d_max = st.number_input(
                "Max \(\\gamma_d\) (\(kN/m^3\))", value=18.5
            )
        rd_calc = RelativeDensityCalculator(
            method="Unit Weight",
            gamma_d=gamma_d,
            gamma_d_min=gamma_d_min,
            gamma_d_max=gamma_d_max,
        )

    dr_val, compactness = rd_calc.compute()
    st.metric("Relative Density (\(D_d\))", f"{dr_val:.1f} %")
    st.info(f"**Compactness State:** `{compactness}`")

st.markdown("---")
if st.button("📥 Export Comprehensive Report (.xlsx)"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inorganic Soils Summary"
    ws["A1"] = "Inorganic Soils Geotechnical Analysis Summary"
    ws["A3"] = "Module"
    ws["B3"] = "Result / Status"
    ws["A4"] = "Visual Description"
    ws["B4"] = final_desc
    ws["A5"] = "Relative Density"
    ws["B5"] = f"{dr_val:.1f}% ({compactness})"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    st.download_button(
        label="Download Formatted Excel Report",
        data=output,
        file_name="Inorganic_Soils_Report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
